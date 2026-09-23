"""Assemble the weekly presentation from the per-person files in slides/.

Each team member edits ONLY their own file in slides/ (e.g. slides/1-ethan-wells.pptx),
so no two people ever touch the same file and git never has a binary merge conflict.
This script stitches those files, in filename order, into the final deck:

    Self-Driving-RC-Car.pptx   (generated -- do not hand-edit; your edits get overwritten)

Run it (from inside the Weekly Reports folder) whenever slides change:

    python3 assemble_deck.py

Naming: the numeric prefix sets slide order (0-title, 1-ethan-wells, ...). Add a member
by dropping a new NN-name.pptx into slides/; remove one by deleting their file.
"""
import copy
import glob
import os
from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.join(HERE, "slides")
OUT = os.path.join(HERE, "Self-Driving-RC-Car.pptx")

# relationships namespace, used to find/rewrite image & hyperlink references (r:embed, r:id, ...)
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def copy_slide(src_slide, dest):
    """Append a deep copy of src_slide (all shapes + images) to dest presentation."""
    # match the source slide's layout by name so placeholders inherit correct position/format
    lname = src_slide.slide_layout.name
    layout = next((l for l in dest.slide_layouts if l.name == lname), dest.slide_layouts[-1])
    new = dest.slides.add_slide(layout)

    # drop the empty placeholders the layout auto-adds; we bring our own shapes over
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)

    # copy every shape's XML
    for sh in src_slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(sh._element))

    # re-create the slide's relationships (images, media, hyperlinks), building a
    # complete old-rId -> new-rId map BEFORE touching the shapes
    rid_map = {}
    for rId, rel in src_slide.part.rels.items():
        if rel.reltype.endswith("slideLayout") or rel.reltype.endswith("notesSlide"):
            continue
        if rel.is_external:
            rid_map[rId] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)

    # then remap in ONE pass, so every r: attribute is rewritten exactly once.
    # Rewriting per relationship instead would let a value written by an earlier
    # mapping (rId1 -> rId2) be matched again by a later one (rId2 -> rId3), which
    # collapses two shapes onto the same image.
    for el in new.shapes._spTree.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(RNS) and val in rid_map:
                el.set(attr, rid_map[val])


def dedupe_media(dest):
    """Give every image/media part a unique name inside the output package.

    Each source deck numbers its own media from image1 up, so two people can both
    ship a ppt/media/image2.jpeg. Those parts keep their original names when they
    are copied in, the saved .pptx then holds two zip entries with the same name,
    and PowerPoint resolves both references to whichever one it reads -- so one
    person's picture silently replaces another's. Renumbering them here keeps each
    picture pointing at its own file.
    """
    used = set()
    n = 0
    for part in dest.part.package.iter_parts():
        name = str(part.partname)
        if not name.startswith("/ppt/media/"):
            continue
        ext = name.rsplit(".", 1)[-1]
        while True:
            n += 1
            new = "/ppt/media/image%d.%s" % (n, ext)
            if new.lower() not in used:
                break
        used.add(new.lower())
        part.partname = PackURI(new)


def main():
    files = sorted(glob.glob(os.path.join(SLIDES_DIR, "*.pptx")))
    if not files:
        raise SystemExit("no slide files found in " + SLIDES_DIR)

    # base the output on the first file so it inherits the theme/layouts, then clear its slides
    dest = Presentation(files[0])
    for sid in list(dest.slides._sldIdLst):
        dest.part.drop_rel(sid.get(qn("r:id")))
        dest.slides._sldIdLst.remove(sid)

    for f in files:
        for slide in Presentation(f).slides:
            copy_slide(slide, dest)

    dedupe_media(dest)

    dest.save(OUT)
    print("assembled %d slides from %d files -> %s"
          % (len(dest.slides._sldIdLst), len(files), os.path.basename(OUT)))


if __name__ == "__main__":
    main()
