import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT = "Self Driving RC Car"
DATE = "10/6/2026"
SUMMARY_DATE = "10/6/2026"

members = ["Ethan Wells (Group Leader)", "Alexis Perez", "Ethan Bishop", "Abigail Duran"]

summary = (
    "This week the team moved from preparing US traffic-sign data to testing a model and narrowing hardware choices for the vehicle. "
    "Ethan Wells completed last week's goal of training and testing a CNN on the prepared 12-class LISA dataset. It reached 89.90% accuracy on 584 held-out test crops. "
    "He reviewed the classes with lower accuracy and prepared 3,449 Mapillary sign crops for the next training run. "
    "Alexis Perez selected and ordered line sensors to help the car follow the course, and completed an OpenCV beginner course. "
    "Ethan Bishop found 3D-printable chassis plate files and continued researching whether ROS 2 would fit the vehicle's software. "
    "Abigail Duran organized the components and connections needed for the custom PCB and reviewed the Compute Module 5 IO Board to determine which features to keep. PCB layout work carries into next week."
)

plan = [
    ("Ethan Wells", "Train and test a CNN on the prepared Mapillary sign data; if time allows, combine the LISA and Mapillary datasets and test a model trained on both."),
    ("Alexis Perez", "Connect the line sensors to the Raspberry Pi and begin using OpenCV with the Raspberry Pi camera."),
    ("Ethan Bishop", "Continue researching ROS 2 and assess the proposed camera, sign detector, behavior, and motor driver nodes for the vehicle."),
    ("Abigail Duran", "Begin the PCB layout, including component placement, board dimensions, power distribution, and signal routing; transfer the selected components into the design."),
]

contributions = [
    ("Ethan Wells", [("10/2/2026", "Trained the CNN on the prepared 12-class LISA sign dataset", 2),
                     ("10/5/2026", "Tested the LISA model on the held-out test set and reviewed the per-class results", 2),
                     ("10/6/2026", "Prepared 3,449 Mapillary sign crops for the same 12 classes", 3)]),
    ("Alexis Perez", [("", "Selected and ordered line sensors for course following", None),
                      ("", "Completed an OpenCV beginner course", None)]),
    ("Ethan Bishop", [("", "Researched 3D-printable chassis plate files", None),
                      ("", "Continued researching a possible ROS 2 implementation for the vehicle", None)]),
    ("Abigail Duran", [("", "Organized the components and connections needed for the custom PCB", None),
                       ("", "Reviewed the Compute Module 5 IO Board to determine which features to keep in the custom design", None)]),
]

hours = [("Ethan Wells", "7", "33"), ("Alexis Perez", "6", "29"),
         ("Ethan Bishop", "6", "28"), ("Abigail Duran", "6", "29")]

d = Document("project_name.docx")

def para_by(pred):
    for p in d.paragraphs:
        if pred(p.text):
            return p
    raise SystemExit("anchor not found: " + repr(pred))

def add_run(p_el, text, bold=False, underline=False):
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    if bold:
        rpr.append(OxmlElement('w:b'))
    if underline:
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rpr.append(sz)   # 12pt
    r.append(rpr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
    p_el.append(r)
    return r

# Use each paragraph type from the TA's template: plain project members,
# first-level bullets for plans and contribution names, and nested bullets for tasks.
p_mem = para_by(lambda t: t.strip() == "Project Members:")
p_sum_head = para_by(lambda t: t.startswith("Weekly Summary"))
p_plan_head = para_by(lambda t: t.startswith("Proposed Plan"))
p_contrib = para_by(lambda t: t.strip() == "Weekly Contributions:")
p_hour = para_by(lambda t: t.strip() == "Hour Tracker:")
member_slot = p_mem._p.getnext()
sum_instr = p_sum_head._p.getnext()
plan_instr = p_plan_head._p.getnext()
member_tmpl = copy.deepcopy(member_slot)
plan_tmpl = copy.deepcopy(plan_instr)
contrib_head_tmpl = copy.deepcopy(p_contrib._p.getnext())
contrib_task_tmpl = copy.deepcopy(p_contrib._p.getnext().getnext())

def fill_para(el, runs):
    """Replace text while preserving the template paragraph's list level and spacing."""
    for r in list(el.findall(qn('w:r'))):
        el.remove(r)
    for text, bold, underline in runs:
        add_run(el, text, bold, underline)
    return el

def clone_para(template, runs):
    return fill_para(copy.deepcopy(template), runs)

# 1) Project Title
p_title = para_by(lambda t: t.startswith("Project Title"))
rs = p_title.runs
rs[0].text = "Project Title: "
rs[1].text = PROJECT
for r in rs[2:]: r.text = ""

# 2) Date
p_date = para_by(lambda t: t.startswith("Date:"))
rs = p_date.runs
rs[0].text = "Date: "
rs[1].text = DATE
for r in rs[2:]: r.text = ""

# 3) Project Members -> fill the plain paragraph, then repeat it
fill_para(member_slot, [(members[0], False, False)])
anchor = member_slot
for m in members[1:]:
    el = clone_para(member_tmpl, [(m, False, False)])
    anchor.addnext(el); anchor = el

# 4) Weekly Summary heading date + body
sum_runs = p_sum_head.runs   # capture once: p_sum_head.runs[0] would be a fresh object each call
for i, r in enumerate(sum_runs):
    r.text = "Weekly Summary (%s):" % SUMMARY_DATE if i == 0 else ""
# replace the template's first-level summary bullet
fill_para(sum_instr, [(summary, False, False)])

# 5) Proposed Plan -> repeat the template's first-level plan bullet
n0, g0 = plan[0]
fill_para(plan_instr, [(n0 + ": ", False, True), (g0, False, False)])
anchor = plan_instr
for n, g in plan[1:]:
    el = clone_para(plan_tmpl, [(n + ": ", False, True), (g, False, False)])
    anchor.addnext(el); anchor = el

# 6) Weekly Contributions -> repeat the template's name/task list levels
start = p_contrib._p
stop = p_hour._p
# collect elements strictly between start and stop
to_remove = []
cur = start.getnext()
while cur is not None and cur is not stop:
    to_remove.append(cur)
    cur = cur.getnext()
for el in to_remove:
    el.getparent().remove(el)
anchor = start
for name, entries in contributions:
    head = clone_para(contrib_head_tmpl, [(name + ":", False, True)])
    anchor.addnext(head); anchor = head
    for date, text, hrs in entries:
        line = ("%s – %s" % (date, text)) if date else text   # date optional
        if hrs is not None:                                   # hours optional (placeholder)
            unit = "hour" if hrs == 1 else "hours"
            line += " (%d %s)" % (hrs, unit)
        el = clone_para(contrib_task_tmpl, [(line, False, False)])
        anchor.addnext(el); anchor = el
# 7) Hour Tracker table
tbl = d.tables[0]
def set_cell(cell, text):
    p = cell.paragraphs[0]
    for r in list(p._p.findall(qn('w:r'))): p._p.remove(r)
    add_run(p._p, text, False, False)
for i, (name, wk, cum) in enumerate(hours):
    row = tbl.rows[i + 1]
    set_cell(row.cells[0], name)
    set_cell(row.cells[1], wk)
    set_cell(row.cells[2], cum)

d.save("Self_Driving_RC_Car.docx")
print("saved")
